import streamlit as st
import json
import io
import os
from PIL import Image
from google import genai
from google.genai import types
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# ==========================================
# 1. הגדרות דף וסגנון RTL לממשק המורה
# ==========================================
st.set_page_config(
    page_title="מחולל מבחנים במתמטיקה | משרד החינוך",
    page_icon="📐",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
    .stApp, .main, [data-testid="stSidebar"] {
        direction: rtl;
        text-align: right;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }
    .stButton>button {
        width: 100%;
        border-radius: 8px;
        height: 3em;
        font-weight: bold;
    }
    .step-card {
        background-color: #f8f9fa;
        border-right: 5px solid #1E88E5;
        padding: 18px;
        border-radius: 6px;
        margin-bottom: 20px;
    }
    .credit-box {
        margin-top: 50px;
        padding: 15px;
        background-color: #e3f2fd;
        border-radius: 8px;
        text-align: center;
        border: 1px solid #bbdefb;
    }
</style>
""", unsafe_allow_html=True)

if "step" not in st.session_state:
    st.session_state.step = 1
if "exam_meta" not in st.session_state:
    st.session_state.exam_meta = {}
if "questions_data" not in st.session_state:
    st.session_state.questions_data = []
if "processed_exam" not in st.session_state:
    st.session_state.processed_exam = None

# ==========================================
# 2. הנדסת פרומפט מערכתי מומחה (Gemini Engine)
# ==========================================
SYSTEM_PROMPT = """
אתה מומחה פדגוגי בכיר למתמטיקה, מתכנן מבחנים ומעריך בחינות בגרות (3, 4 ו-5 יחידות לימוד) במערכת החינוך בישראל.
מטרתך: לקבל תמונות של שאלות מבחן ולשחזר אותן ברמה הגבוהה ביותר עבור:
1. קובץ בחינה לתלמיד (כולל נוסחאות מדויקות ב-LaTeX תקני וסעיפים מוגדרים).
2. קובץ פתרונות מלאים ומפורטים צעד-אחר-צעד, בליווי הסברים מתמטיים מנומקים ותשובות סופיות מודגשות.
3. קובץ מחוון בדיקה אקדמי:
   - פירוק לכל שלב אלגברי/גאומטרי.
   - אחוז הניקוד המדויק לכל שלב (סך האחוזים בכל שאלה חייב להיות בדיוק 100%).
   - קריטריון לביצוע מלא, חלקי, ואפס נקודות.
   - מדיניות טעות נגררת: הגנה מלאה על תלמיד שטעה חישובית והמשיך בדרך נכונה ועקבית.
   - קטלוג טעויות אופייניות לאותו סעיף (הבחנה בין טעות מהותית לטכנית והניקוד שיש להוריד).

הנחיות חמורות:
- אין לנחש! אם פרט כלשהו בתמונה אינו קריא, סמן אותו במפורש כ-[דרוש אימות מורה].
- כל הנוסחאות, המשתנים והביטויים חייבים להיות עטופים ב-LaTeX תקני ($ עבור inline, $$ עבור בלוק).
- בשרטוטי צירים או צורות, יש לתאר את הרכיבים במדויק (שמות קודקודים, זוויות, אורכים ומשוואות).
- הפלט חייב להיות בפורמט JSON בלבד לפי המבנה המוגדר. המפתחות ב-JSON תמיד יהיו באנגלית.
"""

JSON_SCHEMA = """
{
  "questions": [
    {
      "question_number": 1,
      "topic": "נושא השאלה (למשל: חקירת פונקציה מעריכית)",
      "points": 25,
      "text": "נוסח השאלה המלא עם נוסחאות ב-LaTeX",
      "has_figure": true,
      "figure_description": "תיאור מדויק של הגרף או הצורה הגאומטרית",
      "sections": [
        {
          "section_id": "א",
          "text": "נוסח סעיף א'",
          "points": 10
        }
      ],
      "solution": {
        "steps": [
          {
            "step_title": "תחום הגדרה",
            "content": "הסבר מילורי ומשוואות ב-LaTeX",
            "final_answer": "x != 0"
          }
        ]
      },
      "rubric": {
        "steps": [
          {
            "stage_desc": "גזירה נכונה של הפונקציה",
            "percentage": 40,
            "full_credit": "נגזרת מדויקת לחלוטין",
            "partial_credit": "טעות נגזרת פנימית קלה: 20%",
            "zero_credit": "אי גזירה של מכפלה",
            "carried_over_error_policy": "אם גזר לא נכון אך השווה לאפס ופתר עקבית - לקבל את המשך הסעיף",
            "common_errors": [
              {
                "error": "שכחת נגזרת פנימית",
                "severity": "טכנית",
                "deduction_percent": 15
              }
            ]
          }
        ]
      }
    }
  ]
}
"""

def analyze_exam_images_with_gemini(api_key, meta_data, questions_images):
    client = genai.Client(api_key=api_key)
    
    # הוספת הנחיית התרגום במקרה שהמורה בחר שפה אחרת
    language = meta_data.get('language', 'עברית')
    translation_instruction = f"\n\n*** הנחיה קריטית - שפת המבחן: {language} ***\n"
    if language != "עברית":
         translation_instruction += f"עליך לתרגם את כל המלל למורה ולתלמיד (נוסח השאלה, הסעיפים, ההוראות, תיאורי הגרפים, הפתרונות, המחוון וכו') לשפה ה{language} ברמה אקדמית גבוהה ומדויקת מתמטית. הקפד לשמור את כל המפתחות של ה-JSON באנגלית בדיוק כפי שהוגדר, אך כל ערכי הטקסט בתוכם חייבים להיות ב{language}."
    else:
         translation_instruction += "השאר את כל הטקסטים בשפה העברית התקנית."

    contents = [SYSTEM_PROMPT + translation_instruction]
    contents.append(f"פרטי המבחן: {json.dumps(meta_data, ensure_ascii=False)}")
    contents.append("להלן תמונות השאלות לפי הסדר. החזר JSON תקני בלבד התואם ל-Schema:\n" + JSON_SCHEMA)
    
    for q_idx, images in enumerate(questions_images):
        contents.append(f"\n--- תמונות עבור שאלה {q_idx + 1} ---")
        for img in images:
            contents.append(img)
            
    response = client.models.generate_content(
        model='gemini-2.5-pro',
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            temperature=0.1
        )
    )
    return json.loads(response.text)

# ==========================================
# 3. מנוע הפקת מסמכי Word (.docx) מקצועיים
# ==========================================
def set_cell_rtl(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBidi = parse_xml(f'<w:tcBidi {nsdecls("w")} w:val="1"/>')
    tcPr.append(tcBidi)

def create_word_document(exam_data, meta_data, doc_type="exam", logo_file=None):
    doc = docx.Document()
    
    # בדיקה האם השפה היא מימין לשמאל (RTL) או משמאל לימין (LTR)
    language = meta_data.get("language", "עברית")
    is_rtl = language in ["עברית", "ערבית"] 
    align_mode = WD_ALIGN_PARAGRAPH.RIGHT if is_rtl else WD_ALIGN_PARAGRAPH.LEFT
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.5)
    header_table.columns[1].width = Inches(2.2)
    
    cell_info = header_table.cell(0, 0)
    if is_rtl: set_cell_rtl(cell_info)
    p_info = cell_info.paragraphs[0]
    p_info.alignment = align_mode
    
    title_run = p_info.add_run(f"{meta_data.get('school_name', 'School')} - {meta_data.get('exam_name', 'Math Exam')}\n")
    title_run.font.bold = True
    title_run.font.size = Pt(14)
    
    sub_run = p_info.add_run(
        f"שכבה: {meta_data.get('grade', '')} | רמה: {meta_data.get('level', '')} | מועד: {meta_data.get('date', '')}\n"
        f"משך הבחינה: {meta_data.get('duration', '90')} דקות | מורה: {meta_data.get('teacher_name', '')}"
    )
    sub_run.font.size = Pt(10)
    
    cell_logo = header_table.cell(0, 1)
    if is_rtl: set_cell_rtl(cell_logo)
    if logo_file:
        try:
            p_logo = cell_logo.paragraphs[0]
            p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT if is_rtl else WD_ALIGN_PARAGRAPH.RIGHT
            p_logo.add_run().add_picture(io.BytesIO(logo_file), width=Inches(1.5))
        except Exception:
            pass

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.add_run("_______________________________________________________________________________")
    
    p_doc_title = doc.add_paragraph()
    p_doc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc_names = {
        "exam": "טופס בחינה לתלמיד" if is_rtl else "Student Exam Form",
        "solution": "קובץ פתרונות מלאים ומנומקים" if is_rtl else "Full Solutions",
        "rubric": "מחוון בדיקה מפורט" if is_rtl else "Grading Rubric"
    }
    r = p_doc_title.add_run(doc_names.get(doc_type, "בחינה"))
    r.font.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(30, 136, 229)
    
    if doc_type == "exam" and meta_data.get("instructions"):
        p_inst_h = doc.add_paragraph()
        p_inst_h.alignment = align_mode
        r_ih = p_inst_h.add_run("הוראות לנבחן / Instructions:" if not is_rtl else "הוראות לנבחן:")
        r_ih.font.bold = True
        
        for inst in meta_data.get("instructions", "").split("\n"):
            if inst.strip():
                p_i = doc.add_paragraph(inst.strip(), style='List Bullet')
                p_i.alignment = align_mode

    for q in exam_data.get("questions", []):
        q_num = q.get("question_number")
        pts = q.get("points")
        
        q_header = doc.add_paragraph()
        q_header.alignment = align_mode
        q_word = "שאלה" if is_rtl else "Question"
        pts_word = "נקודות" if is_rtl else "points"
        q_run = q_header.add_run(f"\n{q_word} {q_num} ({pts} {pts_word}) - {q.get('topic', '')}")
        q_run.font.bold = True
        q_run.font.size = Pt(13)
        
        if doc_type == "exam":
            p_text = doc.add_paragraph(q.get("text", ""))
            p_text.alignment = align_mode
            for sec in q.get("sections", []):
                p_sec = doc.add_paragraph(f"{sec.get('section_id')}. ({sec.get('points')} {pts_word}) {sec.get('text')}")
                p_sec.alignment = align_mode
                p_sec.paragraph_format.left_indent = Inches(0.2) if not is_rtl else 0
                p_sec.paragraph_format.right_indent = Inches(0.2) if is_rtl else 0
        
        elif doc_type == "solution":
            sol = q.get("solution", {})
            for step in sol.get("steps", []):
                p_st = doc.add_paragraph()
                p_st.alignment = align_mode
                r_st = p_st.add_run(f"• {step.get('step_title')}: ")
                r_st.font.bold = True
                ans_word = "תשובה סופית" if is_rtl else "Final Answer"
                p_st.add_run(f"{step.get('content')}\n{ans_word}: {step.get('final_answer')}")
        
        elif doc_type == "rubric":
            rub = q.get("rubric", {})
            table = doc.add_table(rows=1, cols=4)
            table.autofit = True
            hdr_cells = table.rows[0].cells
            
            headers_rtl = ["שלב בפתרון", "ניקוד ואחוזים", "קריטריון לניקוד חלקי", "טעות נגררת ומדיניות"]
            headers_ltr = ["Step", "Points & %", "Partial Credit", "Carried Error Policy"]
            headers = headers_rtl if is_rtl else headers_ltr
            
            for i, h_text in enumerate(headers):
                if is_rtl: set_cell_rtl(hdr_cells[i])
                p = hdr_cells[i].paragraphs[0]
                p.alignment = align_mode
                p.add_run(h_text).font.bold = True
                
            for step in rub.get("steps", []):
                row_cells = table.add_row().cells
                for cell in row_cells:
                    if is_rtl: set_cell_rtl(cell)
                row_cells[0].paragraphs[0].add_run(step.get("stage_desc", ""))
                row_cells[1].paragraphs[0].add_run(f"{step.get('percentage', 0)}% ({int(pts * step.get('percentage', 0) / 100)})")
                
                partial_txt = f"מלא: {step.get('full_credit')}\nחלקי: {step.get('partial_credit')}" if is_rtl else f"Full: {step.get('full_credit')}\nPartial: {step.get('partial_credit')}"
                row_cells[2].paragraphs[0].add_run(partial_txt)
                
                err_txt = f"נגררת: {step.get('carried_over_error_policy', '')}" if is_rtl else f"Policy: {step.get('carried_over_error_policy', '')}"
                row_cells[3].paragraphs[0].add_run(err_txt)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

# ==========================================
# 4. ממשק האשף ותפריט הצד
# ==========================================
st.sidebar.title("📐 מחולל המבחנים")
steps = [
    "1. הגדרות ופרטי בחינה",
    "2. העלאת שאלוני מקור (תמונות)",
    "3. זיהוי, עיבוד וסנכרון מתמטי",
    "4. תצוגה מקדימה ובדיקות עריכה",
    "5. הפקת קובצי מבחן, פתרון ומחוון"
]
for idx, s in enumerate(steps, 1):
    if st.session_state.step == idx:
        st.sidebar.markdown(f"**🔵 {s}**")
    elif st.session_state.step > idx:
        st.sidebar.markdown(f"✅ {s}")
    else:
        st.sidebar.markdown(f"⚪ {s}")

# --- קרדיט למפתח ---
st.sidebar.markdown("""
<div class="credit-box">
    <b>פותח ע"י שלמה גבר</b><br>
    רכז מתמטיקה, מקיף י'
</div>
""", unsafe_allow_html=True)

# -------------------------------------------------------------
# שלב 1: הגדרות בסיס
# -------------------------------------------------------------
if st.session_state.step == 1:
    st.title("שלב 1: הגדרת הבחינה ופרטי מוסד הלימוד")
    
    col1, col2 = st.columns(2)
    with col1:
        school_name = st.text_input("שם בית הספר", value=st.session_state.exam_meta.get("school_name", "מקיף י'"))
        exam_name = st.text_input("שם הבחינה / נושא", value=st.session_state.exam_meta.get("exam_name", "מבחן מחצית א' במתמטיקה"))
        grade = st.selectbox("שכבת גיל", ["שכבה י'", "שכבה י\"א", "שכבה י\"ב", "חטיבת ביניים"], index=1)
        level = st.selectbox("רמת לימוד", ["5 יחידות לימוד", "4 יחידות לימוד", "3 יחידות לימוד", "רמה מוגברת"], index=0)
    with col2:
        teacher_name = st.text_input("שם המורה / רכז המקצוע", value=st.session_state.exam_meta.get("teacher_name", ""))
        duration = st.number_input("משך הבחינה (בדקות)", min_value=45, max_value=240, value=90, step=15)
        num_questions = st.number_input("מספר שאלות בבחינה", min_value=1, max_value=10, value=3)
        # הוספת שפת הבחינה
        language = st.selectbox("🌍 שפת המבחן (מומלץ לתלמידים עולים חדשים)", 
                                ["עברית", "רוסית", "אנגלית", "צרפתית", "ספרדית", "אמהרית", "אוקראינית"], index=0)

    instructions = st.text_area("הוראות כלליות לתלמיד", value="""1. יש להקפיד על סדר וניקיון בפתרון.
2. כל שלב חישובי או גאומטרי חייב להיות מנומק.
3. טעות חישוב לא תגרור פסילת כל הסעיף אם הדרך הייתה נכונה.
4. שימוש במחשבון מותר בהתאם לתקנון.""", height=120)

    logo_file = st.file_uploader("העלאת סמל/לוגו בית הספר (PNG/JPG)", type=["png", "jpg", "jpeg"])

    if st.button("המשך לשלב העלאת שאלות ⬅️"):
        st.session_state.exam_meta = {
            "school_name": school_name,
            "exam_name": exam_name,
            "grade": grade,
            "level": level,
            "teacher_name": teacher_name,
            "duration": duration,
            "num_questions": num_questions,
            "language": language,
            "instructions": instructions,
            "logo_bytes": logo_file.getvalue() if logo_file else None
        }
        st.session_state.step = 2
        st.rerun()

# -------------------------------------------------------------
# שלב 2: העלאת שאלות
# -------------------------------------------------------------
elif st.session_state.step == 2:
    st.title("שלב 2: העלאת צילומי השאלות")
    
    num_q = st.session_state.exam_meta.get("num_questions", 3)
    uploaded_by_q = []
    
    for i in range(num_q):
        with st.expander(f"📌 אזור העלאה עבור שאלה מס' {i+1}", expanded=(i==0)):
            pts = st.number_input(f"ניקוד עבור שאלה {i+1}", min_value=5, max_value=100, value=100//num_q, key=f"pts_{i}")
            files = st.file_uploader(f"גרור תמונות לשאלה {i+1}", type=["png", "jpg", "jpeg"], accept_multiple_files=True, key=f"upload_{i}")
            uploaded_by_q.append({"question_number": i+1, "points": pts, "files": files})
            
    c1, c2 = st.columns(2)
    with c1:
        if st.button("חזור להגדרות ➡️"):
            st.session_state.step = 1
            st.rerun()
    with c2:
        if st.button("סיום העלאה וסריקה אוטומטית ⬅️"):
            st.session_state.questions_data = uploaded_by_q
            st.session_state.step = 3
            st.rerun()

# -------------------------------------------------------------
# שלב 3: הפעלה מול Gemini API
# -------------------------------------------------------------
elif st.session_state.step == 3:
    st.title("שלב 3: פיענוח מתמטי וסנכרון")
    
    # הודעה מותאמת אם נבחרה שפה זרה
    lang = st.session_state.exam_meta.get("language", "עברית")
    if lang != "עברית":
        st.info(f"🌍 שים לב: המערכת תתרגם אוטומטית את כל השאלות וההוראות לשפה ה{lang}.")

    api_key = st.text_input("הזן מפתח Gemini API של המורה (לצורך העיבוד):", type="password")
    
    if api_key:
        with st.spinner(f"מנתח שאלות ומתרגם ל{lang}..."):
            try:
                all_images_grouped = []
                for q in st.session_state.questions_data:
                    q_imgs = []
                    for f in q["files"]:
                        q_imgs.append(Image.open(f))
                    all_images_grouped.append(q_imgs)
                
                result = analyze_exam_images_with_gemini(api_key, st.session_state.exam_meta, all_images_grouped)
                st.session_state.processed_exam = result
                st.success("הפענוח והתרגום הושלמו בהצלחה!")
                st.session_state.step = 4
                st.rerun()
            except Exception as e:
                st.error(f"שגיאה: {str(e)}")
    else:
        st.info("אנא הזן מפתח API כדי להמשיך.")

# -------------------------------------------------------------
# שלב 4: עריכה ותצוגה מקדימה
# -------------------------------------------------------------
elif st.session_state.step == 4:
    st.title("שלב 4: תצוגה מקדימה ועריכה")
    exam = st.session_state.processed_exam
    
    for i, q in enumerate(exam.get("questions", [])):
        with st.expander(f"📝 שאלה {q.get('question_number')}: {q.get('topic')}", expanded=True):
            q["text"] = st.text_area(f"נוסח שאלה {i+1}", value=q.get("text", ""), height=100, key=f"edit_q_{i}")
            q["points"] = st.number_input(f"ניקוד", value=q.get("points", 25), key=f"edit_pts_{i}")

    c1, c2 = st.columns(2)
    with c1:
        if st.button("חזור ➡️"):
            st.session_state.step = 2
            st.rerun()
    with c2:
        if st.button("אישור והפקת קבצים ⬅️"):
            st.session_state.step = 5
            st.rerun()

# -------------------------------------------------------------
# שלב 5: הפקת קבצים
# -------------------------------------------------------------
elif st.session_state.step == 5:
    st.title("שלב 5: הקבצים מוכנים להורדה!")
    exam_data = st.session_state.processed_exam
    meta_data = st.session_state.exam_meta
    
    docx_exam = create_word_document(exam_data, meta_data, doc_type="exam", logo_file=meta_data.get("logo_bytes"))
    docx_solution = create_word_document(exam_data, meta_data, doc_type="solution")
    docx_rubric = create_word_document(exam_data, meta_data, doc_type="rubric")

    col1, col2, col3 = st.columns(3)
    with col1:
        st.download_button("הורד מבחן (.docx)", data=docx_exam, file_name="מבחן.docx")
    with col2:
        st.download_button("הורד פתרונות (.docx)", data=docx_solution, file_name="פתרונות.docx")
    with col3:
        st.download_button("הורד מחוון (.docx)", data=docx_rubric, file_name="מחוון.docx")

    if st.button("צור מבחן חדש 🔄"):
        st.session_state.step = 1
        st.rerun()